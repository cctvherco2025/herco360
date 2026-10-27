"""Reportes e indicadores de mantenimiento (fase 4).

Para quien planifica (admin de Activos: todo; responsable de categoría: solo
los activos de sus categorías). Filtros comunes: desde/hasta (fecha de la
tarea, YYYY-MM-DD), sucursal, categoría del activo y planificación.

- GET /mantenimiento/reportes          datos del periodo: resumen, por
  planificación / responsable / sucursal, detalle de tareas y fallas.
- GET /mantenimiento/reportes/excel    lo mismo en Excel (Resumen, Detalle
  y Fallas con la foto de cada falla).
- GET /mantenimiento/indicadores       para las gráficas: cumplimiento por
  mes, vencidas por sucursal y fallas por activo, categoría y pregunta.

Cumplimiento = tareas hechas (a tiempo o tarde) / tareas cerradas (hechas +
incompletas); las abiertas y las canceladas no cuentan.
"""
import io
import logging
from datetime import date
from typing import Optional

from fastapi import APIRouter, HTTPException, Depends
from fastapi.responses import StreamingResponse

from core import db, serialize_doc, get_current_user, departamento_de, es_super_admin
import organizacion
import mantenimiento as mant
import routes_activos
import storage

router = APIRouter(prefix='/mantenimiento', tags=['mantenimiento'])
logger = logging.getLogger('mantenimiento')

MESES = ['ene', 'feb', 'mar', 'abr', 'may', 'jun', 'jul', 'ago', 'sep', 'oct', 'nov', 'dic']
ESTADO_TXT = {'completada': 'Hecha', 'incompleta': 'Incompleta', 'pendiente': 'Pendiente', 'programada': 'Programada',
              'cancelada': 'Cancelada', 'vencida': 'Vencida', 'tardia': 'Hecha tarde'}
FOTO_MAX = (160, 120)


def _rango(desde, hasta):
    try:
        d0, d1 = date.fromisoformat(desde), date.fromisoformat(hasta)
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail='Elige el periodo (desde y hasta)')
    if d1 < d0:
        raise HTTPException(status_code=400, detail='La fecha final es anterior a la inicial')
    if (d1 - d0).days > 731:
        raise HTTPException(status_code=400, detail='El periodo puede ser de hasta 2 años')
    return d0, d1


async def _tareas(user, desde, hasta, sucursal=None, categoria=None, plan_id=None, con_respuestas=False):
    await mant._requiere_planificador(user)
    admin, editables = await mant._contexto(user)
    # alcance: las categorías que maneja (quien administra, las de su departamento)
    q = {'fecha': {'$gte': desde, '$lte': hasta}}
    if not es_super_admin(user):
        q['activo_categoria_id'] = {'$in': list(editables)}
    if categoria:
        q['activo_categoria_id'] = categoria if categoria in editables else '__ninguna__'
    if sucursal:
        q['sucursal'] = sucursal
    if plan_id:
        q['plan_id'] = plan_id
    proj = {'_id': 0} if con_respuestas else {'_id': 0, 'respuestas': 0, 'antes': 0, 'despues': 0, 'checklist.items': 0}
    return await db.mant_tareas.find(q, proj).sort([('fecha', 1), ('numero', 1)]).to_list(20000)


def _estado(t, ahora):
    """Estado para el reporte: hecha a tiempo/tarde, incompleta, vencida (abierta y pasó su plazo), pendiente…"""
    if t['estado'] == 'completada':
        return 'tardia' if t.get('tardia') else 'completada'
    if t['estado'] == 'pendiente' and t['vence_en'] <= ahora:
        return 'vencida'
    return t['estado']


def _contar(tareas, ahora):
    c = {'total': 0, 'a_tiempo': 0, 'tardias': 0, 'incompletas': 0, 'abiertas': 0, 'vencidas': 0, 'canceladas': 0, 'fallas': 0}
    for t in tareas:
        e = _estado(t, ahora)
        if e == 'cancelada':
            c['canceladas'] += 1
            continue
        c['total'] += 1
        c['fallas'] += t.get('fallas') or 0
        if e == 'completada':
            c['a_tiempo'] += 1
        elif e == 'tardia':
            c['tardias'] += 1
        elif e == 'incompleta':
            c['incompletas'] += 1
        else:
            c['abiertas'] += 1
            c['vencidas'] += e == 'vencida'
    cerradas = c['a_tiempo'] + c['tardias'] + c['incompletas']
    c['cerradas'] = cerradas
    c['cumplimiento'] = round((c['a_tiempo'] + c['tardias']) / cerradas * 100, 1) if cerradas else None
    return c


def _agrupar(tareas, clave, nombre, ahora):
    grupos = {}
    for t in tareas:
        grupos.setdefault(clave(t), []).append(t)
    out = [{'id': k, 'nombre': nombre(k, ts), **_contar(ts, ahora)} for k, ts in grupos.items()]
    return sorted(out, key=lambda x: (x['cumplimiento'] is None, x['cumplimiento'] if x['cumplimiento'] is not None else 0, x['nombre'] or ''))


def _fallas(tareas):
    """Cada respuesta con falla de las tareas hechas (con respuestas y checklist)."""
    out = []
    for t in tareas:
        if t['estado'] != 'completada' or not t.get('tiene_falla'):
            continue
        items = {i['id']: i for i in t['checklist'].get('items', [])}
        for r in t.get('respuestas', []):
            it = items.get(r['item_id'])
            if not it or not mant._es_falla(it, r.get('valor')):
                continue
            v = r.get('valor')
            respuesta = 'Falla' if it['tipo'] == 'estado' else (', '.join(v) if isinstance(v, list) else str(v))
            out.append({'tarea_id': t['id'], 'numero': t['numero'], 'fecha': t['fecha'], 'completada_en': t.get('completada_en'),
                        'plan_numero': t['plan_numero'], 'activo_id': t['activo_id'], 'activo_codigo': t['activo_codigo'],
                        'activo_nombre': t['activo_nombre'], 'activo_categoria_id': t.get('activo_categoria_id'),
                        'sucursal': t['sucursal'], 'checklist': t['checklist']['nombre'], 'pregunta': it['texto'],
                        'respuesta': respuesta, 'comentario': r.get('comentario') or '', 'fotos': [p['id'] for p in r.get('fotos', [])],
                        'responsable_id': t['responsable_id'], 'responsable': t.get('completada_por_nombre') or ''})
    return out


async def _datos(user, desde, hasta, sucursal, categoria, plan_id):
    _rango(desde, hasta)
    tareas = await _tareas(user, desde, hasta, sucursal, categoria, plan_id, con_respuestas=True)
    ahora = mant._ahora()
    gente = await mant._nombres([t['responsable_id'] for t in tareas])
    cats = {c['id']: c['nombre'] for c in await routes_activos._categorias()}
    detalle = []
    for t in tareas:
        e = _estado(t, ahora)
        fila = {k: t.get(k) for k in ('id', 'numero', 'plan_id', 'plan_numero', 'plan_nombre', 'ciclo_numero', 'activo_codigo',
                                      'activo_nombre', 'sucursal', 'fecha', 'vence_en', 'completada_en', 'fallas')}
        detalle.append({**fila, 'checklist': t['checklist']['nombre'], 'estado': e, 'estado_txt': ESTADO_TXT.get(e, e),
                        'categoria': cats.get(t.get('activo_categoria_id'), ''),
                        'responsable': (gente.get(t['responsable_id']) or {}).get('name', '—')})
    return {
        'desde': desde, 'hasta': hasta, 'ahora': ahora,
        'resumen': _contar(tareas, ahora),
        'por_plan': _agrupar(tareas, lambda t: t['plan_id'], lambda k, ts: f'{ts[0]["plan_numero"]} · {ts[0]["plan_nombre"]}', ahora),
        'por_responsable': _agrupar(tareas, lambda t: t['responsable_id'], lambda k, ts: (gente.get(k) or {}).get('name', '—'), ahora),
        'por_sucursal': _agrupar(tareas, lambda t: t['sucursal'], lambda k, ts: k, ahora),
        'por_categoria': _agrupar(tareas, lambda t: t.get('activo_categoria_id'), lambda k, ts: cats.get(k, '—'), ahora),
        'tareas': detalle,
        'fallas': _fallas(tareas),
    }


@router.get('/reportes')
async def reportes(desde: str, hasta: str, sucursal: Optional[str] = None, categoria: Optional[str] = None,
                   plan_id: Optional[str] = None, user=Depends(get_current_user)):
    return serialize_doc(await _datos(user, desde, hasta, sucursal, categoria, plan_id))


@router.get('/reportes/filtros')
async def filtros(user=Depends(get_current_user)):
    """Opciones de los filtros: sucursales, categorías y planificaciones que la persona ve."""
    await mant._requiere_planificador(user)
    admin, editables = await mant._contexto(user)
    cats = [c for c in await routes_activos._categorias(user) if c['id'] in editables]
    q = {} if es_super_admin(user) else {'categoria_ids': {'$in': list(editables)}}
    planes = await db.mant_planes.find(q, {'_id': 0, 'id': 1, 'numero': 1, 'nombre': 1}).sort('numero', 1).to_list(1000)
    if es_super_admin(user):
        sucursales = [u for d in await organizacion.departamentos() for u in await routes_activos.ubicaciones(d['id'])]
    else:
        sucursales = await routes_activos.ubicaciones(departamento_de(user))
    return serialize_doc({'sucursales': sucursales, 'categorias': [{'id': c['id'], 'nombre': c['nombre']} for c in cats],
                          'planes': planes})


# ── Indicadores ──────────────────────────────────────────────────
@router.get('/indicadores')
async def indicadores(desde: str, hasta: str, sucursal: Optional[str] = None, categoria: Optional[str] = None,
                      plan_id: Optional[str] = None, user=Depends(get_current_user)):
    d0, d1 = _rango(desde, hasta)
    tareas = await _tareas(user, desde, hasta, sucursal, categoria, plan_id, con_respuestas=True)
    ahora = mant._ahora()
    # cumplimiento por mes (los meses del periodo, hasta 12: los últimos)
    meses = []
    y, m = d0.year, d0.month
    while (y, m) <= (d1.year, d1.month):
        meses.append(f'{y}-{m:02d}')
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    meses = meses[-12:]
    por_mes = []
    for ym in meses:
        c = _contar([t for t in tareas if t['fecha'][:7] == ym], ahora)
        por_mes.append({'mes': ym, 'etiqueta': f'{MESES[int(ym[5:]) - 1]} {ym[2:4]}', **c})
    # vencidas e incompletas por sucursal
    suc = {}
    for t in tareas:
        e = _estado(t, ahora)
        if e in ('vencida', 'incompleta'):
            s = suc.setdefault(t['sucursal'], {'sucursal': t['sucursal'], 'vencidas': 0, 'incompletas': 0})
            s['vencidas' if e == 'vencida' else 'incompletas'] += 1
    # fallas
    fallas = _fallas(tareas)
    cats = {c['id']: c['nombre'] for c in await routes_activos._categorias()}

    def top(clave, nombre, n=10):
        cuenta = {}
        for f in fallas:
            k = clave(f)
            cuenta.setdefault(k, {'nombre': nombre(f), 'fallas': 0})
            cuenta[k]['fallas'] += 1
        return sorted(cuenta.values(), key=lambda x: -x['fallas'])[:n]
    return serialize_doc({
        'resumen': _contar(tareas, ahora),
        'por_mes': por_mes,
        'por_sucursal': sorted(suc.values(), key=lambda x: x['sucursal']),
        'fallas_por_activo': top(lambda f: f['activo_id'], lambda f: f'{f["activo_codigo"]} · {f["activo_nombre"]}'),
        'fallas_por_categoria': top(lambda f: f['activo_categoria_id'], lambda f: cats.get(f['activo_categoria_id'], '—')),
        'fallas_por_pregunta': top(lambda f: f['pregunta'].lower(), lambda f: f['pregunta']),
    })


# ── Excel ────────────────────────────────────────────────────────
async def _thumb(path):
    from PIL import Image as PILImage, ImageOps
    try:
        content, _ = await storage.get_object(path)
        im = ImageOps.exif_transpose(PILImage.open(io.BytesIO(content))).convert('RGB')
        im.thumbnail(FOTO_MAX)
        buf = io.BytesIO()
        im.save(buf, 'JPEG', quality=80)
        return buf.getvalue(), im.size
    except Exception as e:
        logger.warning(f'thumb failed for {path}: {e}')
        return None


def _dmy(f):
    return f'{f[8:10]}/{f[5:7]}/{f[0:4]}' if f and len(f) >= 10 else (f or '')


@router.get('/reportes/excel')
async def reportes_excel(desde: str, hasta: str, sucursal: Optional[str] = None, categoria: Optional[str] = None,
                         plan_id: Optional[str] = None, user=Depends(get_current_user)):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
    from openpyxl.drawing.xdr import XDRPositiveSize2D
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.units import pixels_to_EMU

    d = await _datos(user, desde, hasta, sucursal, categoria, plan_id)
    navy, gris = PatternFill('solid', fgColor='1E395E'), PatternFill('solid', fgColor='F1F3F7')
    blanco, negrita = Font(color='FFFFFF', bold=True), Font(bold=True)
    envolver = Alignment(wrap_text=True, vertical='top')

    def encabezado(ws, fila, titulos, anchos=None):
        for i, t in enumerate(titulos, start=1):
            c = ws.cell(row=fila, column=i, value=t)
            c.fill, c.font = navy, blanco
            c.alignment = Alignment(vertical='center', wrap_text=True)
        for i, w in enumerate(anchos or [], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    def pct(v):
        return None if v is None else v / 100

    wb = Workbook()
    # Resumen
    ws = wb.active
    ws.title = 'Resumen'
    ws['A1'] = 'Reporte de mantenimiento preventivo'
    ws['A1'].font = Font(bold=True, size=14, color='1E395E')
    ws['A2'] = f'Periodo: {_dmy(desde)} al {_dmy(hasta)}'
    r = d['resumen']
    filas = [('Tareas del periodo', r['total']), ('Hechas a tiempo', r['a_tiempo']), ('Hechas tarde', r['tardias']),
             ('Incompletas', r['incompletas']), ('Abiertas (vencidas)', f'{r["abiertas"]} ({r["vencidas"]})'),
             ('Fallas reportadas', r['fallas']), ('Cumplimiento', pct(r['cumplimiento']))]
    for i, (k, v) in enumerate(filas, start=4):
        ws.cell(row=i, column=1, value=k).font = negrita
        c = ws.cell(row=i, column=2, value=v)
        if k == 'Cumplimiento':
            c.number_format = '0.0%'
    fila = 4 + len(filas) + 1
    cols = ['Tareas', 'A tiempo', 'Tarde', 'Incompletas', 'Abiertas', 'Fallas', 'Cumplimiento']
    for titulo, grupo in (('Por planificación', 'por_plan'), ('Por responsable', 'por_responsable'),
                          ('Por sucursal', 'por_sucursal'), ('Por categoría', 'por_categoria')):
        ws.cell(row=fila, column=1, value=titulo).font = Font(bold=True, size=12, color='1E395E')
        fila += 1
        encabezado(ws, fila, [titulo.replace('Por ', '').capitalize()] + cols)
        for g in d[grupo]:
            fila += 1
            vals = [g['nombre'], g['total'], g['a_tiempo'], g['tardias'], g['incompletas'], g['abiertas'], g['fallas'], pct(g['cumplimiento'])]
            for j, v in enumerate(vals, start=1):
                c = ws.cell(row=fila, column=j, value=v)
                if j == 8:
                    c.number_format = '0.0%'
        fila += 2
    for i, w in enumerate([38, 10, 10, 10, 12, 10, 10, 14], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Detalle
    ws = wb.create_sheet('Detalle')
    encabezado(ws, 1, ['Tarea', 'Planificación', 'Activo', 'Categoría', 'Sucursal', 'Checklist', 'Responsable', 'Fecha',
                       'Vence', 'Hecha el', 'Estado', 'Fallas'],
               [16, 26, 30, 16, 9, 30, 24, 11, 17, 17, 14, 8])
    for i, t in enumerate(d['tareas'], start=2):
        vals = [t['numero'], f'{t["plan_numero"]} · {t["plan_nombre"]}', f'{t["activo_codigo"]} · {t["activo_nombre"]}', t['categoria'],
                t['sucursal'], t['checklist'], t['responsable'], _dmy(t['fecha']), f'{_dmy(t["vence_en"])} {t["vence_en"][11:16]}',
                f'{_dmy(t["completada_en"])} {(t["completada_en"] or "")[11:16]}'.strip(), t['estado_txt'], t.get('fallas') or 0]
        for j, v in enumerate(vals, start=1):
            ws.cell(row=i, column=j, value=v)
        if i % 2:
            for j in range(1, 13):
                ws.cell(row=i, column=j).fill = gris
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:L{max(1, len(d["tareas"]) + 1)}'

    # Fallas (con la primera foto de cada una)
    ws = wb.create_sheet('Fallas')
    encabezado(ws, 1, ['Foto', 'Fecha', 'Tarea', 'Activo', 'Sucursal', 'Checklist', 'Pregunta', 'Respuesta', 'Descripción', 'Reportó'],
               [24, 11, 16, 30, 9, 26, 32, 18, 40, 22])
    tareas_por_id = {}
    if d['fallas']:
        ids = list({f['tarea_id'] for f in d['fallas']})
        for t in await db.mant_tareas.find({'id': {'$in': ids}}, {'_id': 0, 'id': 1, 'respuestas': 1}).to_list(len(ids)):
            tareas_por_id[t['id']] = {p['id']: p for r in t.get('respuestas', []) for p in r.get('fotos', [])}
    for i, f in enumerate(d['fallas'], start=2):
        vals = [None, _dmy(f['fecha']), f['numero'], f'{f["activo_codigo"]} · {f["activo_nombre"]}', f['sucursal'], f['checklist'],
                f['pregunta'], f['respuesta'], f['comentario'], f['responsable']]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.alignment = envolver
        ws.row_dimensions[i].height = FOTO_MAX[1] * 0.78 + 8 if f['fotos'] else 30
        foto = tareas_por_id.get(f['tarea_id'], {}).get(f['fotos'][0]) if f['fotos'] else None
        t = await _thumb(foto['path']) if foto else None
        if t:
            data, (w, h) = t
            img = XLImage(io.BytesIO(data))
            img.width, img.height = w, h
            img.anchor = OneCellAnchor(_from=AnchorMarker(col=0, row=i - 1, colOff=pixels_to_EMU(4), rowOff=pixels_to_EMU(4)),
                                       ext=XDRPositiveSize2D(pixels_to_EMU(w), pixels_to_EMU(h)))
            ws.add_image(img)
    if not d['fallas']:
        ws.cell(row=2, column=2, value='Sin fallas en el periodo')
    ws.freeze_panes = 'A2'

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nombre = f'Mantenimiento_{desde}_a_{hasta}.xlsx'
    return StreamingResponse(buf, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                             headers={'Content-Disposition': f'attachment; filename="{nombre}"'})
