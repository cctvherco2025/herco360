"""Auditoría de Etiquetas de Precios — ECCP (Experiencia al Cliente y
Control de Pérdidas).

Auditoría semanal: un Coordinador ECCP (o Jefe ECCP, o quien tenga el
permiso 'etiquetas' desde Organigrama) elige a mano sucursal y departamento
y revisa N etiquetas (N configurable, 20 por defecto). Cada etiqueta lleva
resultado (RESULTADOS) y fotos de dos tipos, como el Excel que se llenaba a
mano: 'etiqueta' (la etiqueta en el anaquel, siempre) y 'sistema' (captura del
precio en portal.herco.hn, obligatoria si el precio está mal; en las demás
irregularidades es evidencia opcional). Si el precio está mal, además precio
exhibido y del sistema (la diferencia la calcula el servidor). Fecha, hora, semana ISO y coordinador los pone el servidor.

Sin seguimiento ni estados en esta versión (pedido explícito).

Permisos (core.py): registrar = can_fill_etiquetas; historial =
can_access_etiquetas (todos ven todo, solo lectura); editar, eliminar,
configurar, estadísticas e informes = can_admin_etiquetas.

Configuración (db.etiquetas_config, documento 'current'): cantidad de
etiquetas por auditoría y los rangos de color del cumplimiento. Cada
auditoría guarda la cantidad con que se hizo; el color se calcula siempre con
los rangos vigentes.

Trazabilidad: ediciones y eliminaciones en db.etiquetas_auditoria_log (las
eliminaciones con copia completa). Fotos en Cloudflare R2 (storage.py).
"""
import asyncio
import io
import json
import logging
from datetime import date, timedelta
from typing import List, Optional

from fastapi import APIRouter, HTTPException, Depends, UploadFile, File, Form
from fastapi.responses import Response
from pydantic import BaseModel

from core import (db, serialize_doc, new_id, now_iso, now_local,
                  require_etiquetas_access, require_etiquetas_fill, require_etiquetas_admin)
from notifications import log_activity
import storage

router = APIRouter(prefix='/etiquetas', tags=['etiquetas'])
logger = logging.getLogger('etiquetas')

SUCURSALES = ['H1', 'H2', 'H4', 'H6']
DEPARTAMENTOS = ['Herramientas', 'Hogar', 'Ferretería', 'Iluminación', 'Pinturas', 'Revestimiento']
RESULTADOS = {
    'correcta': 'Correcta',
    'precio_incorrecto': 'Precio incorrecto',
    'sin_etiqueta': 'Sin etiqueta',
    'deteriorada': 'Etiqueta deteriorada o ilegible',
}
DEFAULT_CONFIG = {'cantidad': 20, 'verde_min': 95, 'amarillo_min': 85}
MAX_PHOTO_SIZE = 8 * 1024 * 1024


# ── Configuración ────────────────────────────────────────────────
async def _config():
    doc = await db.etiquetas_config.find_one({'id': 'current'}, {'_id': 0})
    return {**DEFAULT_CONFIG, **{k: doc[k] for k in DEFAULT_CONFIG if doc and k in doc}}


class ConfigUpdate(BaseModel):
    cantidad: int
    verde_min: float
    amarillo_min: float


@router.get('/config')
async def get_config(user=Depends(require_etiquetas_access)):
    return {**await _config(), 'sucursales': SUCURSALES, 'departamentos': DEPARTAMENTOS, 'resultados': RESULTADOS}


@router.put('/config')
async def put_config(data: ConfigUpdate, user=Depends(require_etiquetas_admin)):
    if not 1 <= data.cantidad <= 100:
        raise HTTPException(status_code=400, detail='La cantidad de etiquetas debe estar entre 1 y 100')
    if not 0 <= data.amarillo_min < data.verde_min <= 100:
        raise HTTPException(status_code=400, detail='Los rangos deben cumplir: 0 ≤ amarillo < verde ≤ 100')
    cfg = {'cantidad': data.cantidad, 'verde_min': data.verde_min, 'amarillo_min': data.amarillo_min}
    await db.etiquetas_config.update_one(
        {'id': 'current'},
        {'$set': {**cfg, 'updated_at': now_iso(), 'updated_by': user['name']}}, upsert=True)
    await db.etiquetas_auditoria_log.insert_one({
        'id': new_id(), 'accion': 'configuracion', 'user_id': user['id'], 'user_name': user['name'],
        'at': now_iso(), 'campos': [f'{data.cantidad} etiquetas · verde ≥ {data.verde_min}% · amarillo ≥ {data.amarillo_min}%'],
    })
    return cfg


# ── Utilidades ───────────────────────────────────────────────────
def _txt(v, limit=1000) -> str:
    return str(v or '').strip()[:limit]


def _precio(v, n, campo):
    if v in (None, ''):
        raise HTTPException(status_code=400, detail=f'Etiqueta {n}: indica el {campo}')
    try:
        x = round(float(v), 2)
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail=f'Etiqueta {n}: el {campo} no es un número válido')
    if x < 0:
        raise HTTPException(status_code=400, detail=f'Etiqueta {n}: el {campo} no puede ser negativo')
    return x


def _clean_etiqueta(e, n):
    resultado = _txt(e.get('resultado'))
    if resultado not in RESULTADOS:
        raise HTTPException(status_code=400, detail=f'Etiqueta {n}: elige el resultado de la verificación')
    out = {
        'numero': n, 'resultado': resultado,
        'precio_exhibido': None, 'precio_sistema': None, 'diferencia': None,
        'observaciones': _txt(e.get('observaciones')) if resultado != 'correcta' else '',
    }
    if resultado == 'precio_incorrecto':
        out['precio_exhibido'] = _precio(e.get('precio_exhibido'), n, 'precio exhibido')
        out['precio_sistema'] = _precio(e.get('precio_sistema'), n, 'precio del sistema')
        out['diferencia'] = round(out['precio_exhibido'] - out['precio_sistema'], 2)
    return out


def _totales(etiquetas):
    c = {k: sum(1 for e in etiquetas if e['resultado'] == k) for k in RESULTADOS}
    total = len(etiquetas)
    return {
        'total': total, 'correctas': c['correcta'], 'precio_incorrecto': c['precio_incorrecto'],
        'sin_etiqueta': c['sin_etiqueta'], 'deterioradas': c['deteriorada'],
        'irregularidades': total - c['correcta'],
        'cumplimiento': round(c['correcta'] / total * 100, 1) if total else 0,
    }


def _semana(d: date):
    y, w, _ = d.isocalendar()
    return f'{y}-W{w:02d}', w


async def _upload(aud_id, f: UploadFile):
    content = await f.read()
    if not content:
        return None
    if len(content) > MAX_PHOTO_SIZE:
        raise HTTPException(status_code=400, detail=f'Una foto supera el máximo de {MAX_PHOTO_SIZE // (1024 * 1024)} MB')
    ctype = f.content_type or 'image/jpeg'
    photo_id = new_id()
    ext = 'png' if 'png' in ctype else ('webp' if 'webp' in ctype else 'jpg')
    path = f'{storage.APP_NAME}/etiquetas/{aud_id}/{photo_id}.{ext}'
    try:
        await storage.put_object(path, content, ctype)
    except Exception as e:
        logger.error(f'photo upload failed: {e}')
        raise HTTPException(status_code=502, detail='No se pudo subir una de las fotos')
    return {'id': photo_id, 'path': path, 'content_type': ctype}


def _filtro(sucursal=None, departamento=None, coordinador=None, desde=None, hasta=None, semana=None, mes=None):
    q = {}
    if sucursal:
        q['sucursal'] = sucursal
    if departamento:
        q['departamento'] = departamento
    if coordinador:
        q['coordinador_id'] = coordinador
    if semana:
        q['semana'] = semana  # 'YYYY-Www'
    if mes:
        q['mes'] = mes  # 'YYYY-MM'
    rango = {}
    if desde:
        rango['$gte'] = desde
    if hasta:
        rango['$lte'] = hasta
    if rango:
        q['fecha_local'] = rango
    return q


async def _log(accion, aud, user, **extra):
    await db.etiquetas_auditoria_log.insert_one({
        'id': new_id(), 'accion': accion, 'auditoria_id': aud['id'],
        'sucursal': aud.get('sucursal'), 'departamento': aud.get('departamento'),
        'coordinador_name': aud.get('coordinador_name'), 'fecha_local': aud.get('fecha_local'),
        'user_id': user['id'], 'user_name': user['name'], 'at': now_iso(), **extra,
    })


# ── Registrar ────────────────────────────────────────────────────
@router.post('/auditorias')
async def create_auditoria(
    data: str = Form(...),
    photos: List[UploadFile] = File(default=[]),
    photo_owner: List[str] = Form(default=[]),
    user=Depends(require_etiquetas_fill),
):
    try:
        payload = json.loads(data)
    except Exception:
        raise HTTPException(status_code=400, detail='Datos de la auditoría inválidos')
    sucursal = _txt(payload.get('sucursal'))
    if sucursal not in SUCURSALES:
        raise HTTPException(status_code=400, detail='Selecciona la sucursal')
    departamento = _txt(payload.get('departamento'))
    if departamento not in DEPARTAMENTOS:
        raise HTTPException(status_code=400, detail='Selecciona el departamento')
    cfg = await _config()
    raw = payload.get('etiquetas') or []
    if len(raw) != cfg['cantidad']:
        raise HTTPException(status_code=400, detail=f'La auditoría debe tener {cfg["cantidad"]} etiquetas revisadas (vienen {len(raw)})')
    if len(photos) != len(photo_owner):
        raise HTTPException(status_code=400, detail='Las fotos no coinciden con sus etiquetas')

    etiquetas, by_key = [], {}
    for n, e in enumerate(raw, start=1):
        item = {'id': new_id(), **_clean_etiqueta(e, n), 'photos': []}
        etiquetas.append(item)
        by_key[_txt(e.get('key'))] = item
    # photo_owner = "<key>|etiqueta" (foto de la etiqueta en el anaquel) o
    # "<key>|sistema" (captura del precio en el sistema / evidencia adicional)
    owners = [(o.rsplit('|', 1) + ['etiqueta'])[:2] for o in photo_owner]
    tiene = {(k, t) for k, t in owners}
    for key, item in by_key.items():
        if (key, 'etiqueta') not in tiene:
            raise HTTPException(status_code=400, detail=f'Etiqueta {item["numero"]}: agrega la foto de la etiqueta')
        if item['resultado'] == 'precio_incorrecto' and (key, 'sistema') not in tiene:
            raise HTTPException(status_code=400, detail=f'Etiqueta {item["numero"]}: agrega la captura del precio en el sistema')

    aud_id = new_id()
    for f, (key, tipo) in zip(photos, owners):
        if key in by_key:
            ph = await _upload(aud_id, f)
            if ph:
                by_key[key]['photos'].append({**ph, 'tipo': 'sistema' if tipo == 'sistema' else 'etiqueta'})

    local = now_local()
    semana, semana_num = _semana(local.date())
    doc = {
        'id': aud_id, 'sucursal': sucursal, 'departamento': departamento,
        'coordinador_id': user['id'], 'coordinador_name': user['name'], 'coordinador_avatar': user.get('avatar_url'),
        'created_at': now_iso(), 'fecha_local': local.strftime('%Y-%m-%d'), 'hora_local': local.strftime('%H:%M'),
        'semana': semana, 'semana_num': semana_num, 'mes': local.strftime('%Y-%m'),
        'cantidad_config': cfg['cantidad'],
        'etiquetas': etiquetas, **_totales(etiquetas), 'cambios': [],
    }
    await db.etiquetas_auditorias.insert_one(doc)
    doc.pop('_id', None)
    await log_activity(user['id'], user['name'], user.get('avatar_url'),
                       'registró una auditoría de etiquetas en', f'{sucursal} · {departamento}', 'etiquetas')
    return serialize_doc(doc)


# ── Historial ────────────────────────────────────────────────────
@router.get('/auditorias')
async def list_auditorias(sucursal: Optional[str] = None, departamento: Optional[str] = None,
                          coordinador: Optional[str] = None, desde: Optional[str] = None,
                          hasta: Optional[str] = None, semana: Optional[str] = None, mes: Optional[str] = None,
                          user=Depends(require_etiquetas_access)):
    rows = await db.etiquetas_auditorias.find(
        _filtro(sucursal, departamento, coordinador, desde, hasta, semana, mes),
        {'_id': 0, 'etiquetas': 0, 'cambios': 0},
    ).sort('created_at', -1).limit(500).to_list(500)
    return serialize_doc(rows)


@router.get('/coordinadores')
async def list_coordinadores(user=Depends(require_etiquetas_access)):
    rows = await db.etiquetas_auditorias.aggregate([
        {'$group': {'_id': '$coordinador_id', 'name': {'$last': '$coordinador_name'}}},
        {'$sort': {'name': 1}},
    ]).to_list(200)
    return [{'id': r['_id'], 'name': r['name']} for r in rows]


@router.get('/auditorias/{aud_id}')
async def get_auditoria(aud_id: str, user=Depends(require_etiquetas_access)):
    row = await db.etiquetas_auditorias.find_one({'id': aud_id}, {'_id': 0})
    if not row:
        raise HTTPException(status_code=404, detail='Auditoría no encontrada')
    return serialize_doc(row)


# ── Editar / eliminar (administrador) ────────────────────────────
class AuditoriaUpdate(BaseModel):
    sucursal: str
    departamento: str
    etiquetas: list  # [{id, resultado, precio_exhibido, precio_sistema, observaciones}]


CAMPOS_LABEL = {
    'resultado': 'resultado', 'precio_exhibido': 'precio exhibido',
    'precio_sistema': 'precio del sistema', 'observaciones': 'observaciones',
}


@router.put('/auditorias/{aud_id}')
async def update_auditoria(aud_id: str, data: AuditoriaUpdate, user=Depends(require_etiquetas_admin)):
    """Corrige sucursal, departamento y los datos de cada etiqueta. No se
    agregan ni quitan etiquetas, y las fotos se conservan tal cual."""
    row = await db.etiquetas_auditorias.find_one({'id': aud_id}, {'_id': 0})
    if not row:
        raise HTTPException(status_code=404, detail='Auditoría no encontrada')
    sucursal, departamento = _txt(data.sucursal), _txt(data.departamento)
    if sucursal not in SUCURSALES:
        raise HTTPException(status_code=400, detail='Selecciona la sucursal')
    if departamento not in DEPARTAMENTOS:
        raise HTTPException(status_code=400, detail='Selecciona el departamento')
    nuevos_by_id = {_txt(e.get('id')): e for e in data.etiquetas}

    cambios, etiquetas = [], []
    if row['sucursal'] != sucursal:
        cambios.append('sucursal')
    if row['departamento'] != departamento:
        cambios.append('departamento')
    for old in row.get('etiquetas', []):
        e = nuevos_by_id.get(old['id'])
        if e is None:
            etiquetas.append(old)
            continue
        clean = _clean_etiqueta(e, old['numero'])
        dif = [lbl for k, lbl in CAMPOS_LABEL.items() if old.get(k) != clean.get(k)]
        if dif:
            cambios.append(f'etiqueta {old["numero"]} ({", ".join(dif)})')
        etiquetas.append({**old, **clean})
    if not cambios:
        return serialize_doc(row)

    cambio = {'user_id': user['id'], 'user_name': user['name'], 'at': now_iso(), 'campos': cambios}
    await db.etiquetas_auditorias.update_one({'id': aud_id}, {
        '$set': {'sucursal': sucursal, 'departamento': departamento, 'etiquetas': etiquetas, **_totales(etiquetas),
                 'updated_at': cambio['at'], 'updated_by': user['name']},
        '$push': {'cambios': cambio},
    })
    await _log('editada', {**row, 'sucursal': sucursal, 'departamento': departamento}, user, campos=cambios)
    out = await db.etiquetas_auditorias.find_one({'id': aud_id}, {'_id': 0})
    return serialize_doc(out)


@router.delete('/auditorias/{aud_id}')
async def delete_auditoria(aud_id: str, user=Depends(require_etiquetas_admin)):
    row = await db.etiquetas_auditorias.find_one({'id': aud_id}, {'_id': 0})
    if not row:
        raise HTTPException(status_code=404, detail='Auditoría no encontrada')
    await _log('eliminada', row, user, snapshot=row)
    await db.etiquetas_auditorias.delete_one({'id': aud_id})
    return {'message': 'Auditoría eliminada'}


@router.get('/log')
async def list_log(user=Depends(require_etiquetas_admin)):
    rows = await db.etiquetas_auditoria_log.find({}, {'_id': 0, 'snapshot': 0}).sort('at', -1).limit(300).to_list(300)
    return serialize_doc(rows)


@router.get('/auditorias/{aud_id}/foto/{photo_id}')
async def get_photo(aud_id: str, photo_id: str, user=Depends(require_etiquetas_access)):
    row = await db.etiquetas_auditorias.find_one({'id': aud_id}, {'_id': 0, 'etiquetas.photos': 1})
    if not row:
        raise HTTPException(status_code=404, detail='Auditoría no encontrada')
    photo = next((p for e in row.get('etiquetas', []) for p in e.get('photos', []) if p['id'] == photo_id), None)
    if not photo:
        raise HTTPException(status_code=404, detail='Foto no encontrada')
    try:
        content, ctype = await storage.get_object(photo['path'])
    except Exception as e:
        logger.error(f'photo download failed: {e}')
        raise HTTPException(status_code=502, detail='No se pudo descargar la foto')
    return Response(content=content, media_type=photo.get('content_type', ctype))


# ── Estadísticas (administrador) ─────────────────────────────────
def _agg(rows):
    total = sum(r.get('total', 0) for r in rows)
    ok = sum(r.get('correctas', 0) for r in rows)
    return {'auditorias': len(rows), 'etiquetas': total, 'correctas': ok, 'irregularidades': total - ok,
            'cumplimiento': round(ok / total * 100, 1) if total else None}


@router.get('/estadisticas')
async def estadisticas(desde: str, hasta: str, sucursal: Optional[str] = None,
                       departamento: Optional[str] = None, user=Depends(require_etiquetas_admin)):
    """Resumen del periodo desde..hasta + comparativos: las 8 semanas ISO y
    los 6 meses que terminan en `hasta` (con los mismos filtros)."""
    fin = date.fromisoformat(hasta)
    lunes_fin = fin - timedelta(days=fin.weekday())
    semanas = [lunes_fin - timedelta(weeks=i) for i in range(7, -1, -1)]
    meses = []
    y, m = fin.year, fin.month
    for _ in range(6):
        meses.insert(0, f'{y}-{m:02d}')
        y, m = (y, m - 1) if m > 1 else (y - 1, 12)
    inicio = min(desde, semanas[0].isoformat(), f'{meses[0]}-01')

    proj = {'_id': 0, 'sucursal': 1, 'departamento': 1, 'fecha_local': 1, 'semana': 1, 'mes': 1,
            'total': 1, 'correctas': 1, 'precio_incorrecto': 1, 'sin_etiqueta': 1, 'deterioradas': 1}
    rows = await db.etiquetas_auditorias.find(
        _filtro(sucursal=sucursal, departamento=departamento, desde=inicio, hasta=hasta), proj).to_list(10000)
    periodo = [r for r in rows if desde <= r['fecha_local'] <= hasta]

    res = _agg(periodo)
    for k in ('precio_incorrecto', 'sin_etiqueta', 'deterioradas'):
        res[k] = sum(r.get(k, 0) for r in periodo)
    return {
        'resumen': res,
        'por_sucursal': [{'sucursal': s, **_agg([r for r in periodo if r['sucursal'] == s])} for s in SUCURSALES],
        'por_departamento': [{'departamento': d, **_agg([r for r in periodo if r['departamento'] == d])} for d in DEPARTAMENTOS],
        'semanal': [{'semana': _semana(l)[0], 'desde': l.isoformat(), **_agg([r for r in rows if r['semana'] == _semana(l)[0]])}
                    for l in semanas],
        'mensual': [{'mes': mm, **_agg([r for r in rows if r['mes'] == mm])} for mm in meses],
    }


# ── Informes (administrador) ─────────────────────────────────────
@router.get('/informe')
async def informe_datos(desde: str, hasta: str, sucursal: Optional[str] = None,
                        departamento: Optional[str] = None, user=Depends(require_etiquetas_admin)):
    """Auditorías completas del periodo (el PDF se arma en el navegador)."""
    rows = await db.etiquetas_auditorias.find(
        _filtro(sucursal=sucursal, departamento=departamento, desde=desde, hasta=hasta), {'_id': 0, 'cambios': 0},
    ).sort('created_at', 1).to_list(1000)
    return serialize_doc(rows)


FOTO_MAX = (240, 300)  # px: así entra en la celda de las hojas por auditoría


async def _thumb(path, size=FOTO_MAX):
    """Foto reducida a JPEG para incrustar en Excel (None si falla)."""
    from PIL import Image as PILImage, ImageOps
    try:
        content, _ = await storage.get_object(path)
        im = ImageOps.exif_transpose(PILImage.open(io.BytesIO(content))).convert('RGB')
        im.thumbnail(size)
        buf = io.BytesIO()
        im.save(buf, 'JPEG', quality=82)
        return buf.getvalue(), im.size
    except Exception as e:
        logger.warning(f'thumb failed for {path}: {e}')
        return None


def _fotos(e, tipo):
    return [p for p in e.get('photos', []) if (p.get('tipo') or 'etiqueta') == tipo]


def _dmy(f):
    return f'{f[8:10]}/{f[5:7]}/{f[0:4]}' if f and len(f) == 10 else (f or '')


@router.get('/informe.xlsx')
async def informe_excel(desde: str, hasta: str, sucursal: Optional[str] = None,
                        departamento: Optional[str] = None, user=Depends(require_etiquetas_admin)):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    cfg = await _config()
    rows = await db.etiquetas_auditorias.find(
        _filtro(sucursal=sucursal, departamento=departamento, desde=desde, hasta=hasta), {'_id': 0, 'cambios': 0},
    ).sort('created_at', 1).to_list(1000)

    sem = asyncio.Semaphore(8)

    async def get_thumb(p):
        async with sem:
            return await _thumb(p['path'])
    # en la hoja va la primera foto de cada tipo; el resto se menciona en Observaciones
    jobs = {}
    for r in rows:
        for e in r.get('etiquetas', []):
            for tipo in ('etiqueta', 'sistema'):
                for p in _fotos(e, tipo)[:1]:
                    if p['id'] not in jobs:
                        jobs[p['id']] = get_thumb(p)
    thumbs = dict(zip(jobs.keys(), await asyncio.gather(*jobs.values())))

    from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
    from openpyxl.drawing.xdr import XDRPositiveSize2D
    from openpyxl.utils.units import pixels_to_EMU

    head = PatternFill('solid', fgColor='1E395E')
    fills = {'g': PatternFill('solid', fgColor='DCFCE7'), 'a': PatternFill('solid', fgColor='FEF9C3'),
             'r': PatternFill('solid', fgColor='FEE2E2')}
    tone = lambda pct: 'g' if pct >= cfg['verde_min'] else ('a' if pct >= cfg['amarillo_min'] else 'r')
    wrap = Alignment(wrap_text=True, vertical='top')
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    money = '#,##0.00'

    def header_row(ws, row, cols):
        for j, h in enumerate(cols, start=1):
            c = ws.cell(row=row, column=j, value=h)
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = head
            c.alignment = center

    wb = Workbook()
    # Hoja 1: informe resumido (una fila por auditoría)
    ws = wb.active
    ws.title = 'Resumen'
    header_row(ws, 1, ['Fecha', 'Semana', 'Sucursal', 'Departamento', 'Coordinador', 'Etiquetas verificadas', 'Correctas',
                       'Precio incorrecto', 'Sin etiqueta', 'Deterioradas', 'Irregularidades', 'Cumplimiento %'])
    for i, w in enumerate([11, 10, 9, 14, 24, 11, 10, 10, 10, 11, 13, 13], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    for r in rows:
        ws.append([_dmy(r['fecha_local']), r.get('semana', ''), r['sucursal'], r['departamento'], r['coordinador_name'],
                   r['total'], r['correctas'], r['precio_incorrecto'], r['sin_etiqueta'], r['deterioradas'],
                   r['irregularidades'], r['cumplimiento']])
        ws.cell(row=ws.max_row, column=12).fill = fills[tone(r['cumplimiento'])]
    if rows:
        tot = sum(r['total'] for r in rows)
        ok = sum(r['correctas'] for r in rows)
        ws.append(['TOTAL', '', '', '', f'{len(rows)} auditorías', tot, ok,
                   sum(r['precio_incorrecto'] for r in rows), sum(r['sin_etiqueta'] for r in rows),
                   sum(r['deterioradas'] for r in rows), tot - ok, round(ok / tot * 100, 1) if tot else 0])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)

    # Una hoja por auditoría, como el Excel que se llenaba a mano: título con
    # tienda y departamento, columnas "Etiqueta" | "Mal etiquetado" con las
    # fotos grandes y, al lado, los datos de cada etiqueta.
    FOTO_COL_W = 36  # ancho de columna (~257 px)
    ROW_PX = FOTO_MAX[1] + 16
    nombres = set()

    def poner_foto(ws, t, col, row):
        data, (w, h) = t
        img = XLImage(io.BytesIO(data))  # un BytesIO por imagen: openpyxl lo lee al guardar
        img.width, img.height = w, h
        off_x = max(0, (FOTO_COL_W * 7 + 5 - w) // 2)  # centrada en la celda
        off_y = max(0, (ROW_PX - h) // 2)
        img.anchor = OneCellAnchor(
            _from=AnchorMarker(col=col, row=row - 1, colOff=pixels_to_EMU(off_x), rowOff=pixels_to_EMU(off_y)),
            ext=XDRPositiveSize2D(pixels_to_EMU(w), pixels_to_EMU(h)))
        ws.add_image(img)

    for r in rows:
        semnum = int(str(r.get('semana') or '-W0').split('-W')[-1] or 0)
        base = f'{r["sucursal"]} {r["departamento"]} S{semnum}'[:28]
        nombre, k = base, 2
        while nombre in nombres:
            nombre, k = f'{base[:25]} ({k})', k + 1
        nombres.add(nombre)
        ws = wb.create_sheet(nombre)
        ws.merge_cells('A1:H1')
        t = ws.cell(row=1, column=1, value=(
            f'Tienda {r["sucursal"]} ({r["departamento"]}) · Semana {semnum} · {_dmy(r["fecha_local"])} · '
            f'{r["coordinador_name"]} · Cumplimiento {r["cumplimiento"]}%'))
        t.font = Font(bold=True, size=13, color='FFFFFF')
        t.fill = head
        t.alignment = Alignment(vertical='center')
        ws.row_dimensions[1].height = 26
        header_row(ws, 2, ['Etiqueta', 'Mal etiquetado', 'N°', 'Resultado', 'Precio exhibido',
                           'Precio sistema', 'Diferencia', 'Observaciones'])
        for col, w in zip('ABCDEFGH', (FOTO_COL_W, FOTO_COL_W, 5, 18, 12, 12, 11, 34)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = 'A3'

        for i, e in enumerate(r.get('etiquetas', [])):
            row = 3 + i
            ws.row_dimensions[row].height = ROW_PX * 0.75
            fe, fs = _fotos(e, 'etiqueta'), _fotos(e, 'sistema')
            obs = e.get('observaciones', '')
            extras = []
            if len(fe) > 1:
                extras.append(f'+{len(fe) - 1} foto(s) de la etiqueta')
            if len(fs) > 1:
                extras.append(f'+{len(fs) - 1} captura(s)/evidencia')
            if extras:
                obs = (obs + '\n' if obs else '') + f'({", ".join(extras)} en HERCO OPS)'
            ws.append([None, None, e['numero'], RESULTADOS.get(e['resultado'], e['resultado']),
                       e.get('precio_exhibido'), e.get('precio_sistema'), e.get('diferencia'), obs])
            for c in ws[row][2:]:
                c.alignment = wrap
            ws.cell(row=row, column=3).font = Font(bold=True)
            res = ws.cell(row=row, column=4)
            res.font = Font(bold=True, color='16A34A' if e['resultado'] == 'correcta' else 'DC2626')
            if e['resultado'] != 'correcta':
                res.fill = fills['r']
            for col in (5, 6, 7):
                ws.cell(row=row, column=col).number_format = money
            if e.get('diferencia'):
                ws.cell(row=row, column=7).font = Font(bold=True, color='DC2626')
            for col, fotos in ((0, fe), (1, fs)):
                if fotos and thumbs.get(fotos[0]['id']):
                    poner_foto(ws, thumbs[fotos[0]['id']], col, row)

    buf = io.BytesIO()
    wb.save(buf)
    extra = ''.join(f'_{x}' for x in (sucursal, departamento) if x)
    fname = f'Auditoria_Etiquetas{extra}_{desde}_a_{hasta}.xlsx'
    return Response(
        content=buf.getvalue(),
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename="{fname}"'},
    )
