"""Recorridos Preventivos en Tienda — ECCP (Experiencia al Cliente y Control
de Pérdidas).

Un Coordinador ECCP (o Jefe ECCP, o quien tenga el permiso 'recorridos' desde
Organigrama) recorre una sucursal (H1/H2/H4/H6, elegida a mano) y registra
cero o más hallazgos. Cada hallazgo: categoría (solo las 6 de CATEGORIAS),
tipo de irregularidad (uno de TIPOS o uno escrito a mano, ver _tipo), producto, cantidad, descripción, al menos
una foto, responsable notificado y acción. Sin hallazgos = "Sin novedades".
Fecha/hora y coordinador los pone el servidor.

No hay estados: el seguimiento lo hace el administrador agregando notas
(acción correctiva + fotos de la solución) a cada hallazgo; cada nota guarda
quién y cuándo.

Permisos (core.py): registrar = can_fill_recorridos; historial =
can_access_recorridos (todos ven todo, solo lectura); editar, eliminar,
seguimiento, estadísticas e informes = can_admin_recorridos.

Trazabilidad: ediciones, seguimientos y eliminaciones quedan en
db.recorridos_auditoria (las eliminaciones con copia completa del recorrido).
Las fotos van a Cloudflare R2 (storage.py) y solo se descargan autenticado.
"""
import asyncio
import io
import json
import logging
from typing import List, Optional

from fastapi import APIRouter, HTTPException, Depends, UploadFile, File, Form
from fastapi.responses import Response
from pydantic import BaseModel

from core import (db, serialize_doc, new_id, now_iso, now_local,
                  require_recorridos_access, require_recorridos_fill, require_recorridos_admin)
from notifications import log_activity
import storage

router = APIRouter(prefix='/recorridos', tags=['recorridos'])
logger = logging.getLogger('recorridos')

SUCURSALES = ['H1', 'H2', 'H4', 'H6']
CATEGORIAS = ['Herramientas', 'Hogar', 'Ferretería', 'Iluminación', 'Pinturas', 'Revestimiento']
TIPOS = [
    'Producto sin etiqueta de precio.',
    'Precio incorrecto o mal colocado.',
    'Dispositivo de seguridad ausente.',
    'Dispositivo de seguridad mal colocado.',
    'Dispositivo de seguridad defectuoso.',
    'Candado en mal estado.',
    'Vitrina o exhibidor sin asegurar.',
    'Basura en góndolas o exhibidores.',
    'Producto fuera de su ubicación.',
    'Pasillos obstruidos.',
    'Equipo de exhibición o publicidad sin funcionar.',
    'Otros.',
]
MAX_PHOTO_SIZE = 8 * 1024 * 1024  # 8 MB por foto
MAX_HALLAZGOS = 60


# ── Utilidades ───────────────────────────────────────────────────
def _txt(v, limit=2000) -> str:
    return str(v or '').strip()[:limit]


def _cantidad(v):
    if v in (None, ''):
        return None
    try:
        n = int(v)
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail='La cantidad debe ser un número entero')
    if n < 0:
        raise HTTPException(status_code=400, detail='La cantidad no puede ser negativa')
    return n


def _clave_tipo(t: str) -> str:
    return ' '.join(t.lower().rstrip('.').split())


async def _tipos_usados() -> list:
    """Tipos escritos a mano en recorridos anteriores (fuera de TIPOS)."""
    usados = await db.recorridos_preventivos.distinct('hallazgos.tipo')
    oficiales = {_clave_tipo(t) for t in TIPOS}
    vistos, out = set(), []
    for t in sorted(usados, key=str.lower):
        k = _clave_tipo(t or '')
        if k and k not in oficiales and k not in vistos:
            vistos.add(k)
            out.append(t)
    return out


def _tipo(raw, conocidos) -> str:
    """Tipo de irregularidad: si coincide (sin importar mayúsculas, espacios
    ni el punto final) con uno oficial o uno ya usado, se guarda ese nombre
    exacto para que las estadísticas no se partan; si no, el texto escrito."""
    t = ' '.join(_txt(raw, 120).split())
    if not t:
        return ''
    k = _clave_tipo(t)
    for c in conocidos:
        if _clave_tipo(c) == k:
            return c
    return t[0].upper() + t[1:]


def _clean_hallazgo(h, n, conocidos=TIPOS):
    """Campos editables de un hallazgo, validados. n = número para el mensaje."""
    categoria = _txt(h.get('categoria'))
    if categoria not in CATEGORIAS:
        raise HTTPException(status_code=400, detail=f'Hallazgo {n}: elige el departamento')
    tipo = _tipo(h.get('tipo'), conocidos)
    if not tipo:
        raise HTTPException(status_code=400, detail=f'Hallazgo {n}: indica el tipo de irregularidad')
    descripcion = _txt(h.get('descripcion'))
    if not descripcion:
        raise HTTPException(status_code=400, detail=f'Hallazgo {n}: escribe la descripción')
    return {
        'categoria': categoria, 'tipo': tipo,
        'producto': _txt(h.get('producto'), 300), 'cantidad': _cantidad(h.get('cantidad')),
        'descripcion': descripcion,
        'responsable': _txt(h.get('responsable'), 200), 'accion': _txt(h.get('accion')),
    }


def _resumen(hallazgos):
    cats = [c for c in CATEGORIAS if any(h['categoria'] == c for h in hallazgos)]
    return {'total_hallazgos': len(hallazgos), 'categorias': cats, 'sin_novedades': len(hallazgos) == 0}


async def _upload(rec_id, f: UploadFile):
    content = await f.read()
    if not content:
        return None
    if len(content) > MAX_PHOTO_SIZE:
        raise HTTPException(status_code=400, detail=f'Una foto supera el máximo de {MAX_PHOTO_SIZE // (1024 * 1024)} MB')
    ctype = f.content_type or 'image/jpeg'
    photo_id = new_id()
    ext = 'png' if 'png' in ctype else ('webp' if 'webp' in ctype else 'jpg')
    path = f'{storage.APP_NAME}/recorridos/{rec_id}/{photo_id}.{ext}'
    try:
        await storage.put_object(path, content, ctype)
    except Exception as e:
        logger.error(f'photo upload failed: {e}')
        raise HTTPException(status_code=502, detail='No se pudo subir una de las fotos')
    return {'id': photo_id, 'path': path, 'content_type': ctype}


def _filtro(sucursal=None, categoria=None, tipo=None, coordinador=None, desde=None, hasta=None):
    q = {}
    if sucursal:
        q['sucursal'] = sucursal
    if categoria:
        q['hallazgos.categoria'] = categoria
    if tipo:
        q['hallazgos.tipo'] = tipo
    if coordinador:
        q['coordinador_id'] = coordinador
    rango = {}
    if desde:
        rango['$gte'] = desde  # 'YYYY-MM-DD' local
    if hasta:
        rango['$lte'] = hasta
    if rango:
        q['fecha_local'] = rango
    return q


async def _audit(accion, rec, user, **extra):
    await db.recorridos_auditoria.insert_one({
        'id': new_id(), 'accion': accion, 'recorrido_id': rec['id'],
        'sucursal': rec.get('sucursal'), 'coordinador_name': rec.get('coordinador_name'),
        'fecha_local': rec.get('fecha_local'),
        'user_id': user['id'], 'user_name': user['name'], 'at': now_iso(), **extra,
    })


# ── Catálogo ─────────────────────────────────────────────────────
@router.get('/catalogo')
async def catalogo(user=Depends(require_recorridos_access)):
    return {'sucursales': SUCURSALES, 'categorias': CATEGORIAS, 'tipos': TIPOS}


@router.get('/tipos')
async def tipos(user=Depends(require_recorridos_access)):
    """Oficiales + los escritos a mano antes (sugerencias y filtro)."""
    return {'oficiales': TIPOS, 'otros': await _tipos_usados()}


# ── Registrar ────────────────────────────────────────────────────
@router.post('/registros')
async def create_recorrido(
    data: str = Form(...),
    photos: List[UploadFile] = File(default=[]),
    photo_owner: List[str] = Form(default=[]),
    user=Depends(require_recorridos_fill),
):
    try:
        payload = json.loads(data)
    except Exception:
        raise HTTPException(status_code=400, detail='Datos del recorrido inválidos')
    sucursal = _txt(payload.get('sucursal'))
    if sucursal not in SUCURSALES:
        raise HTTPException(status_code=400, detail='Selecciona la sucursal')
    raw = payload.get('hallazgos') or []
    if len(raw) > MAX_HALLAZGOS:
        raise HTTPException(status_code=400, detail=f'Máximo {MAX_HALLAZGOS} hallazgos por recorrido')
    if len(photos) != len(photo_owner):
        raise HTTPException(status_code=400, detail='Las fotos no coinciden con sus hallazgos')

    conocidos = TIPOS + await _tipos_usados()
    hallazgos, by_key = [], {}
    for n, h in enumerate(raw, start=1):
        item = {'id': new_id(), **_clean_hallazgo(h, n, conocidos), 'photos': [], 'seguimiento': []}
        hallazgos.append(item)
        by_key[_txt(h.get('key'))] = (n, item)
    owners = set(photo_owner)
    for key, (n, _) in by_key.items():
        if key not in owners:
            raise HTTPException(status_code=400, detail=f'Hallazgo {n}: agrega al menos una foto')

    rec_id = new_id()
    for f, owner in zip(photos, photo_owner):
        if owner not in by_key:
            continue
        ph = await _upload(rec_id, f)
        if ph:
            by_key[owner][1]['photos'].append(ph)

    local = now_local()
    doc = {
        'id': rec_id, 'sucursal': sucursal,
        'coordinador_id': user['id'], 'coordinador_name': user['name'], 'coordinador_avatar': user.get('avatar_url'),
        'created_at': now_iso(), 'fecha_local': local.strftime('%Y-%m-%d'), 'hora_local': local.strftime('%H:%M'),
        'hallazgos': hallazgos, **_resumen(hallazgos), 'cambios': [],
    }
    await db.recorridos_preventivos.insert_one(doc)
    doc.pop('_id', None)
    await log_activity(user['id'], user['name'], user.get('avatar_url'),
                       'registró un recorrido preventivo en', sucursal, 'recorridos')
    return serialize_doc(doc)


# ── Historial ────────────────────────────────────────────────────
@router.get('/registros')
async def list_recorridos(sucursal: Optional[str] = None, categoria: Optional[str] = None,
                          tipo: Optional[str] = None, coordinador: Optional[str] = None,
                          desde: Optional[str] = None, hasta: Optional[str] = None,
                          user=Depends(require_recorridos_access)):
    rows = await db.recorridos_preventivos.find(
        _filtro(sucursal, categoria, tipo, coordinador, desde, hasta),
        {'_id': 0, 'hallazgos': 0, 'cambios': 0},
    ).sort('created_at', -1).limit(500).to_list(500)
    return serialize_doc(rows)


@router.get('/coordinadores')
async def list_coordinadores(user=Depends(require_recorridos_access)):
    """Quienes han registrado recorridos (para el filtro del historial)."""
    rows = await db.recorridos_preventivos.aggregate([
        {'$group': {'_id': '$coordinador_id', 'name': {'$last': '$coordinador_name'}}},
        {'$sort': {'name': 1}},
    ]).to_list(200)
    return [{'id': r['_id'], 'name': r['name']} for r in rows]


@router.get('/registros/{rec_id}')
async def get_recorrido(rec_id: str, user=Depends(require_recorridos_access)):
    row = await db.recorridos_preventivos.find_one({'id': rec_id}, {'_id': 0})
    if not row:
        raise HTTPException(status_code=404, detail='Recorrido no encontrado')
    return serialize_doc(row)


# ── Editar / eliminar (administrador) ────────────────────────────
class RecorridoUpdate(BaseModel):
    sucursal: str
    hallazgos: list  # [{id, categoria, tipo, producto, cantidad, descripcion, responsable, accion}]


CAMPOS_LABEL = {
    'categoria': 'departamento', 'tipo': 'tipo', 'producto': 'producto', 'cantidad': 'cantidad',
    'descripcion': 'descripción', 'responsable': 'responsable', 'accion': 'acción',
}


@router.put('/registros/{rec_id}')
async def update_recorrido(rec_id: str, data: RecorridoUpdate, user=Depends(require_recorridos_admin)):
    """Edita sucursal y los campos de texto de los hallazgos existentes; un
    hallazgo que no venga en la lista se quita. Fotos y seguimiento se
    conservan tal cual."""
    row = await db.recorridos_preventivos.find_one({'id': rec_id}, {'_id': 0})
    if not row:
        raise HTTPException(status_code=404, detail='Recorrido no encontrado')
    sucursal = _txt(data.sucursal)
    if sucursal not in SUCURSALES:
        raise HTTPException(status_code=400, detail='Selecciona la sucursal')

    prev = {h['id']: h for h in row.get('hallazgos', [])}
    conocidos = TIPOS + await _tipos_usados()
    cambios, nuevos = [], []
    if row.get('sucursal') != sucursal:
        cambios.append('sucursal')
    for n, h in enumerate(data.hallazgos, start=1):
        old = prev.get(_txt(h.get('id')))
        if not old:
            continue  # aquí no se agregan hallazgos nuevos
        clean = _clean_hallazgo(h, n, conocidos)
        dif = [CAMPOS_LABEL[k] for k, v in clean.items() if old.get(k) != v]
        if dif:
            cambios.append(f'hallazgo {n} ({", ".join(dif)})')
        nuevos.append({**old, **clean})
    quitados = [h for hid, h in prev.items() if hid not in {x['id'] for x in nuevos}]
    for h in quitados:
        cambios.append(f'quitó hallazgo "{h["tipo"].rstrip(".")}" de {h["categoria"]}')
    if not cambios:
        return serialize_doc(row)

    cambio = {'user_id': user['id'], 'user_name': user['name'], 'at': now_iso(), 'campos': cambios}
    await db.recorridos_preventivos.update_one({'id': rec_id}, {
        '$set': {'sucursal': sucursal, 'hallazgos': nuevos, **_resumen(nuevos),
                 'updated_at': cambio['at'], 'updated_by': user['name']},
        '$push': {'cambios': cambio},
    })
    await _audit('editado', {**row, 'sucursal': sucursal}, user, campos=cambios,
                 **({'hallazgos_quitados': quitados} if quitados else {}))
    out = await db.recorridos_preventivos.find_one({'id': rec_id}, {'_id': 0})
    return serialize_doc(out)


@router.delete('/registros/{rec_id}')
async def delete_recorrido(rec_id: str, user=Depends(require_recorridos_admin)):
    row = await db.recorridos_preventivos.find_one({'id': rec_id}, {'_id': 0})
    if not row:
        raise HTTPException(status_code=404, detail='Recorrido no encontrado')
    await _audit('eliminado', row, user, snapshot=row)
    await db.recorridos_preventivos.delete_one({'id': rec_id})
    return {'message': 'Recorrido eliminado'}


# ── Seguimiento (administrador) ──────────────────────────────────
@router.post('/registros/{rec_id}/hallazgos/{hid}/seguimiento')
async def add_seguimiento(
    rec_id: str, hid: str,
    data: str = Form(...),
    photos: List[UploadFile] = File(default=[]),
    user=Depends(require_recorridos_admin),
):
    try:
        payload = json.loads(data)
    except Exception:
        raise HTTPException(status_code=400, detail='Datos del seguimiento inválidos')
    accion = _txt(payload.get('accion'))
    if not accion:
        raise HTTPException(status_code=400, detail='Escribe la acción correctiva')
    row = await db.recorridos_preventivos.find_one({'id': rec_id}, {'_id': 0})
    if not row:
        raise HTTPException(status_code=404, detail='Recorrido no encontrado')
    hallazgo = next((h for h in row.get('hallazgos', []) if h['id'] == hid), None)
    if not hallazgo:
        raise HTTPException(status_code=404, detail='Hallazgo no encontrado')

    fotos = []
    for f in photos:
        ph = await _upload(rec_id, f)
        if ph:
            fotos.append(ph)
    nota = {'id': new_id(), 'accion': accion, 'photos': fotos,
            'user_id': user['id'], 'user_name': user['name'], 'at': now_iso()}
    await db.recorridos_preventivos.update_one(
        {'id': rec_id, 'hallazgos.id': hid}, {'$push': {'hallazgos.$.seguimiento': nota}})
    await _audit('seguimiento', row, user,
                 campos=[f'{hallazgo["categoria"]} · {hallazgo["tipo"].rstrip(".")}: {accion[:120]}'])
    out = await db.recorridos_preventivos.find_one({'id': rec_id}, {'_id': 0})
    return serialize_doc(out)


@router.get('/auditoria')
async def list_auditoria(user=Depends(require_recorridos_admin)):
    rows = await db.recorridos_auditoria.find(
        {}, {'_id': 0, 'snapshot': 0, 'hallazgos_quitados': 0}
    ).sort('at', -1).limit(300).to_list(300)
    return serialize_doc(rows)


# ── Fotos ────────────────────────────────────────────────────────
def _all_photos(row):
    for h in row.get('hallazgos', []):
        yield from h.get('photos', [])
        for s in h.get('seguimiento', []):
            yield from s.get('photos', [])


@router.get('/registros/{rec_id}/foto/{photo_id}')
async def get_photo(rec_id: str, photo_id: str, user=Depends(require_recorridos_access)):
    row = await db.recorridos_preventivos.find_one({'id': rec_id}, {'_id': 0})
    if not row:
        raise HTTPException(status_code=404, detail='Recorrido no encontrado')
    photo = next((p for p in _all_photos(row) if p['id'] == photo_id), None)
    if not photo:
        raise HTTPException(status_code=404, detail='Foto no encontrada')
    try:
        content, ctype = await storage.get_object(photo['path'])
    except Exception as e:
        logger.error(f'photo download failed: {e}')
        raise HTTPException(status_code=502, detail='No se pudo descargar la foto')
    return Response(content=content, media_type=photo.get('content_type', ctype))


# ── Estadísticas e informes (administrador) ──────────────────────
@router.get('/estadisticas')
async def estadisticas(desde: Optional[str] = None, hasta: Optional[str] = None,
                       sucursal: Optional[str] = None, user=Depends(require_recorridos_admin)):
    rows = await db.recorridos_preventivos.find(
        _filtro(sucursal=sucursal, desde=desde, hasta=hasta),
        {'_id': 0, 'sucursal': 1, 'hallazgos.categoria': 1, 'hallazgos.tipo': 1, 'hallazgos.seguimiento.id': 1},
    ).to_list(5000)
    por_sucursal = {s: 0 for s in SUCURSALES}
    por_categoria = {c: 0 for c in CATEGORIAS}
    por_tipo = {t: 0 for t in TIPOS}
    total_h = con_seg = sin_nov = 0
    for r in rows:
        hs = r.get('hallazgos') or []
        if not hs:
            sin_nov += 1
        for h in hs:
            total_h += 1
            por_sucursal[r['sucursal']] = por_sucursal.get(r['sucursal'], 0) + 1
            por_categoria[h['categoria']] = por_categoria.get(h['categoria'], 0) + 1
            por_tipo[h['tipo']] = por_tipo.get(h['tipo'], 0) + 1
            if h.get('seguimiento'):
                con_seg += 1
    desc = lambda d, key: sorted(({key: k, 'n': v} for k, v in d.items()), key=lambda x: -x['n'])
    return {
        'total_recorridos': len(rows), 'recorridos_sin_novedades': sin_nov,
        'total_hallazgos': total_h, 'hallazgos_con_seguimiento': con_seg,
        'por_sucursal': [{'sucursal': s, 'n': por_sucursal[s]} for s in SUCURSALES],
        'por_categoria': desc(por_categoria, 'categoria'),
        'por_tipo': desc(por_tipo, 'tipo'),
    }


@router.get('/informe')
async def informe_datos(desde: str, hasta: str, sucursal: Optional[str] = None,
                        user=Depends(require_recorridos_admin)):
    """Recorridos completos del periodo (para armar el PDF en el navegador)."""
    rows = await db.recorridos_preventivos.find(
        _filtro(sucursal=sucursal, desde=desde, hasta=hasta), {'_id': 0, 'cambios': 0},
    ).sort('created_at', 1).to_list(2000)
    return serialize_doc(rows)


def _fmt_fecha(rec):
    f = rec.get('fecha_local') or ''
    return f'{f[8:10]}/{f[5:7]}/{f[0:4]}' if len(f) == 10 else f


def _seguimiento_txt(h):
    out = []
    for s in h.get('seguimiento') or []:
        at = (s.get('at') or '')[:10]
        out.append(f'{at[8:10]}/{at[5:7]}/{at[0:4]} {s.get("user_name", "")}: {s.get("accion", "")}')
    return '\n'.join(out)


THUMB_PX = 110
MAX_FOTOS_XLSX = 4


async def _thumb(path):
    """Miniatura JPEG para incrustar en Excel (None si falla)."""
    from PIL import Image as PILImage
    try:
        content, _ = await storage.get_object(path)
        im = PILImage.open(io.BytesIO(content))
        im = im.convert('RGB')
        im.thumbnail((THUMB_PX, THUMB_PX))
        buf = io.BytesIO()
        im.save(buf, 'JPEG', quality=80)
        buf.seek(0)
        return buf, im.size
    except Exception as e:
        logger.warning(f'thumb failed for {path}: {e}')
        return None


@router.get('/informe.xlsx')
async def informe_excel(desde: str, hasta: str, sucursal: Optional[str] = None,
                        user=Depends(require_recorridos_admin)):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    rows = await db.recorridos_preventivos.find(
        _filtro(sucursal=sucursal, desde=desde, hasta=hasta), {'_id': 0, 'cambios': 0},
    ).sort('created_at', 1).to_list(2000)

    # miniaturas (máx. MAX_FOTOS_XLSX por hallazgo), en paralelo y con límite
    sem = asyncio.Semaphore(8)

    async def get_thumb(p):
        async with sem:
            return await _thumb(p['path'])
    jobs = {}
    for r in rows:
        for h in r.get('hallazgos', []):
            for p in h.get('photos', [])[:MAX_FOTOS_XLSX]:
                jobs[p['id']] = get_thumb(p)
    thumbs = dict(zip(jobs.keys(), await asyncio.gather(*jobs.values())))

    head = PatternFill('solid', fgColor='1E395E')
    wrap = Alignment(wrap_text=True, vertical='top')
    wb = Workbook()

    ws = wb.active
    ws.title = 'Hallazgos'
    cols = ['Fecha', 'Hora', 'Sucursal', 'Coordinador', 'Departamento', 'Tipo de irregularidad', 'Producto',
            'Cantidad', 'Descripción', 'Responsable notificado', 'Acción realizada o recomendada', 'Seguimiento',
            'Total de fotos']
    # el Excel muestra hasta MAX_FOTOS_XLSX miniaturas; el total va en su columna
    foto_cols = [f'Foto {i + 1}' for i in range(MAX_FOTOS_XLSX)]
    ws.append(cols + foto_cols)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = head
        c.alignment = Alignment(vertical='center', wrap_text=True)
    widths = [11, 7, 9, 24, 14, 30, 22, 9, 40, 22, 36, 36, 9] + [16] * MAX_FOTOS_XLSX
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'

    first_photo_col = len(cols) + 1
    row_i = 1
    for r in rows:
        for h in r.get('hallazgos', []):
            row_i += 1
            ws.append([
                _fmt_fecha(r), r.get('hora_local', ''), r.get('sucursal', ''), r.get('coordinador_name', ''),
                h.get('categoria', ''), (h.get('tipo') or '').rstrip('.'), h.get('producto', ''),
                h.get('cantidad'), h.get('descripcion', ''), h.get('responsable', ''), h.get('accion', ''),
                _seguimiento_txt(h), len(h.get('photos', [])),
            ])
            for c in ws[row_i]:
                c.alignment = wrap
            fotos = h.get('photos', [])
            placed = 0
            for p in fotos[:MAX_FOTOS_XLSX]:
                t = thumbs.get(p['id'])
                if not t:
                    continue
                buf, (w, hpx) = t
                img = XLImage(buf)
                img.width, img.height = w, hpx
                ws.add_image(img, f'{get_column_letter(first_photo_col + placed)}{row_i}')
                placed += 1
            ws.row_dimensions[row_i].height = 88 if placed else 45

    ws2 = wb.create_sheet('Recorridos')
    ws2.append(['Fecha', 'Hora', 'Sucursal', 'Coordinador', 'Departamentos inspeccionados', 'Total de hallazgos', 'Resultado'])
    for c in ws2[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = head
    for r in rows:
        ws2.append([_fmt_fecha(r), r.get('hora_local', ''), r.get('sucursal', ''), r.get('coordinador_name', ''),
                    ', '.join(r.get('categorias') or []) or '—', r.get('total_hallazgos', 0),
                    'Sin novedades' if r.get('sin_novedades') else 'Con hallazgos'])
    for col, w in zip('ABCDEFG', (11, 7, 9, 24, 40, 16, 16)):
        ws2.column_dimensions[col].width = w

    buf = io.BytesIO()
    wb.save(buf)
    suc = f'_{sucursal}' if sucursal else ''
    fname = f'Recorridos_Preventivos{suc}_{desde}_a_{hasta}.xlsx'
    return Response(
        content=buf.getvalue(),
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename="{fname}"'},
    )
